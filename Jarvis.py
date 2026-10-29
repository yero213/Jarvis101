"""J.A.R.V.I.S. - lokale spraakassistent (gewone versie)."""
import csv
import io
import os
import queue
import re
import shutil
import subprocess
import sys
import threading
import time
import urllib.request
from collections import deque
from pathlib import Path


def _ollama_beschikbaar() -> bool:
    """Controleer of Ollama draait op localhost:11434."""
    try:
        req = urllib.request.Request("http://localhost:11434/api/tags", method="GET")
        with urllib.request.urlopen(req, timeout=2) as resp:
            return resp.status == 200
    except Exception:
        return False


if __name__ == "__main__" and not _ollama_beschikbaar():
    print("Ollama niet gevonden op localhost:11434.")
    print("Schakel over op Google AI Studio (Gemini)...")
    try:
        import numpy as np   # noqa: F401
        import sounddevice as sd   # noqa: F401
        from dotenv import load_dotenv   # noqa: F401
        from Jarvis_gemini import run_session
        import asyncio
        asyncio.run(run_session())
    except ImportError as e:
        print(f"Kan Gemini niet starten: {e}")
        print("Installeer: py -m pip install google-genai python-dotenv numpy sounddevice")
        sys.exit(1)
    except Exception as e:
        print(f"Fout bij Gemini-opstart: {e}")
        sys.exit(1)
    sys.exit(0)

# Ollama is beschikbaar → laad alle lokale dependencies
import numpy as np
import sounddevice as sd
from faster_whisper import WhisperModel
from ollama import Client
from kokoro_onnx import Kokoro


# --- .env laden (zonder harde dependency op python-dotenv) ---
try:
    from dotenv import load_dotenv  # type: ignore
    load_dotenv()
except ImportError:
    _env_file = Path(__file__).with_name(".env")
    if _env_file.is_file():
        for _line in _env_file.read_text(encoding="utf-8").splitlines():
            _line = _line.strip()
            if _line and not _line.startswith("#") and "=" in _line:
                _k, _v = _line.split("=", 1)
                os.environ.setdefault(_k.strip(), _v.strip().strip("\"'"))

# --- Instellingen ---
LLM_MODEL = "qwen2.5-coder:7b"
WHISPER_MODEL = "small"      # "base" is sneller, "small"/"medium" nauwkeuriger
LANGUAGE = "en"              # "nl" als je Nederlands praat, of None voor auto-detectie
WAKE_WORDS = ["jarvis", "jervis", "jarvus", "jorvis", "jarvish", "parvis", "parvish", "gervis", "george"]
THRESHOLD = 0.015            # verhoog bij achtergrondlawaai, verlaag als hij je niet hoort
SILENCE_SECONDS = 1.0        # stilte waarna de opname stopt (lager = sneller antwoord, maar hij kan je afkappen)
MAX_SECONDS = 15
SLEEP_WORDS = ["go to sleep", "stand by", "standby"]
IDLE_TIMEOUT = 180           # seconden zonder gesprek waarna hij weer een wake word nodig heeft (0 = nooit)

# --- Ollama ---
KEEP_ALIVE = "30m"           # houdt het model in het geheugen, zodat de eerste vraag na een pauze niet traag is
NUM_CTX = 8192               # contextgrootte van het model
MAX_HISTORY_CHARS = 12000    # daarboven worden de oudste berichten uit het gesprek gehaald

# --- Audio ---
# Microfoon en speakers draaien op DEZELFDE samplerate en de streams blijven de hele tijd open.
# Steeds openen/sluiten of wisselen van samplerate geeft op een Mac hoorbare tikken.
AUDIO_RATE = int(sd.query_devices(kind="output")["default_samplerate"])
WHISPER_RATE = 16000         # Whisper wil 16 kHz; de opname wordt daarvoor omgerekend
BLOCK = AUDIO_RATE // 16     # ~62 ms per blok

# --- Stem-instellingen ---
VOICE_A = "bm_george"
VOICE_B = "bm_fable"         # probeer ook "bm_lewis" of "bm_daniel"
VOICE_MIX = 0.6              # 0.6 = 60% A, 40% B
SPEED = 0.95                 # iets langzamer klinkt natuurlijker
PAUSE = 0.15                 # seconden stilte tussen zinnen

# --- Onderbreken ---
BARGE_IN = True              # je kunt Jarvis onderbreken terwijl hij praat
BARGE_THRESHOLD = 0.04       # aanhoudend microfoonniveau dat als "jij praat" telt (moet boven zijn eigen echo liggen)
BARGE_BLOCKS = 4             # zoveel opeenvolgende luide blokken (~0,25 s) nodig om te onderbreken
STOP_WORDS = ["stop", "quiet", "enough", "shut up", "hold on", "wait"]

# --- Bestanden ---
WORKSPACE = Path.home() / "Jarvis" / "workspace"   # Jarvis mag ALLEEN in deze map lezen en schrijven
MAX_READ_CHARS = 6000        # meer tekst dan dit wordt afgekapt (houdt het model snel en op de rails)
MAX_TOOL_ROUNDS = 5          # maximaal aantal bestandsacties per vraag

SYSTEM_PROMPT = (
    "You are J.A.R.V.I.S., a sophisticated, dry-witted British AI assistant. "
    "Your replies are spoken aloud, so talk like a person, not a document: "
    "short sentences, natural phrasing, contractions where fitting, "
    "no lists, no markdown, no emojis, no code unless asked. "
    "Keep answers to one to three sentences unless more is needed. "
    "Address the user as Sir. "
    "You can list, read, write and append to files in your workspace using the provided tools. "
    "If you are unsure of a file name, call list_files first. "
    "Never claim you read or wrote a file unless the tool result confirms it. "
    "When you summarise a file, speak the gist in a few sentences instead of reading it out."
)

# --- Modellen laden ---
print("Modellen laden...")
stt = WhisperModel(WHISPER_MODEL, device="cpu", compute_type="int8")
ollama_client = Client(host="http://localhost:11434")
kokoro = Kokoro("kokoro-v1.0.onnx", "voices-v1.0.bin")

# Gemengde stem: een gemiddelde van twee stemprofielen klinkt vaak minder "synthetisch"
voice = (
    kokoro.get_voice_style(VOICE_A) * VOICE_MIX
    + kokoro.get_voice_style(VOICE_B) * (1 - VOICE_MIX)
)

history = [{"role": "system", "content": SYSTEM_PROMPT}]


# --- Audio: microfoon en speaker ---
def resample(x: np.ndarray, src: int, dst: int) -> np.ndarray:
    """Rekent audio om naar een andere samplerate."""
    x = np.asarray(x, dtype=np.float32)
    if src == dst or len(x) == 0:
        return x
    if src % dst == 0:                       # bijv. 48000 -> 16000: middelen per groepje
        k = src // dst
        n = len(x) // k * k
        return x[:n].reshape(-1, k).mean(axis=1).astype(np.float32)
    n = int(len(x) * dst / src)              # bijv. 24000 -> 48000
    return np.interp(np.linspace(0, len(x) - 1, n), np.arange(len(x)), x).astype(np.float32)


def drain(q: queue.Queue):
    while True:
        try:
            q.get_nowait()
        except queue.Empty:
            break


mic_q = queue.Queue()


def _mic_callback(indata, frames, time_info, status):
    mic_q.put(indata[:, 0].copy())


mic_stream = sd.InputStream(samplerate=AUDIO_RATE, channels=1, dtype="float32",
                            blocksize=BLOCK, callback=_mic_callback)
mic_stream.start()


class Speaker:
    """Eén doorlopende uitvoerstream met zachte in- en uitfades: geen tikken bij zinsovergangen of afbreken."""
    FADE = 0.008   # seconden

    def __init__(self, rate: int):
        self.rate = rate
        self.lock = threading.Lock()
        self.buf = np.zeros(0, dtype=np.float32)
        self.gen = 0   # verhoogt bij elke nieuwe uitspraak of stop, zodat oude zinnen niet meer binnenkomen
        self.stream = sd.OutputStream(samplerate=rate, channels=1, dtype="float32",
                                      callback=self._callback)
        self.stream.start()

    def _callback(self, outdata, frames, time_info, status):
        with self.lock:
            n = min(frames, len(self.buf))
            outdata[:n, 0] = self.buf[:n]
            outdata[n:, 0] = 0.0
            self.buf = self.buf[n:]

    def begin(self) -> int:
        with self.lock:
            self.gen += 1
            self.buf = self.buf[:0]
            return self.gen

    def add(self, samples: np.ndarray, gen: int):
        samples = np.asarray(samples, dtype=np.float32).copy()
        n = min(int(self.rate * self.FADE), len(samples) // 2)
        if n > 0:
            ramp = np.linspace(0.0, 1.0, n, dtype=np.float32)
            samples[:n] *= ramp
            samples[-n:] *= ramp[::-1]
        with self.lock:
            if gen == self.gen:
                self.buf = np.concatenate([self.buf, samples])

    def pending(self) -> int:
        with self.lock:
            return len(self.buf)

    def stop(self):
        """Kapt het afspelen af met een korte fade-out."""
        with self.lock:
            self.gen += 1
            n = min(int(self.rate * 0.02), len(self.buf))
            self.buf = self.buf[:n] * np.linspace(1.0, 0.0, n, dtype=np.float32) if n else self.buf[:0]


speaker = Speaker(AUDIO_RATE)


def _record(q: queue.Queue, frames: list | None = None) -> np.ndarray:
    """Leest microfoonblokken uit q. Begint bij spraak (of meteen als frames al gevuld zijn,
    bijvoorbeeld na een onderbreking) en stopt na een moment van stilte."""
    pre_roll = deque(maxlen=5)  # bewaart de laatste blokken zodat het begin niet wordt afgeknipt
    frames = list(frames) if frames else []
    started = bool(frames)
    silent = 0.0
    block_dur = BLOCK / AUDIO_RATE

    while True:
        block = q.get()
        rms = float(np.sqrt(np.mean(block ** 2)))

        if not started:
            pre_roll.append(block)
            if rms > THRESHOLD:
                started = True
                frames.extend(pre_roll)
            continue

        frames.append(block)
        silent = 0.0 if rms > THRESHOLD else silent + block_dur
        if silent >= SILENCE_SECONDS or len(frames) * block_dur >= MAX_SECONDS:
            break

    return np.concatenate(frames).flatten()


def listen() -> np.ndarray:
    """Wacht op spraak, neem op tot er stilte valt, en geef de audio terug."""
    drain(mic_q)
    return _record(mic_q)


def transcribe(audio: np.ndarray) -> str:
    audio = resample(audio, AUDIO_RATE, WHISPER_RATE)
    segments, _ = stt.transcribe(audio, language=LANGUAGE, beam_size=1, vad_filter=True)
    text = " ".join(s.text.strip() for s in segments).strip()
    return text


# --- Bestandstools ---
WORKSPACE.mkdir(parents=True, exist_ok=True)
TEXT_LIKE = {".txt", ".md", ".csv", ".json", ".py", ".html", ".xml", ".yaml", ".yml", ".log", ".ini", ".toml", ".js", ".css", ".sh", ".tsv"}


def _safe_path(name: str) -> Path:
    """Geeft een pad terug binnen de workspace, of gooit een fout (ook bij ../ trucs)."""
    root = WORKSPACE.resolve()
    p = (root / name).resolve()
    if p != root and root not in p.parents:
        raise ValueError("That path is outside the workspace.")
    return p


def _truncate(text: str) -> str:
    if len(text) > MAX_READ_CHARS:
        return text[:MAX_READ_CHARS] + f"\n[...truncated, file has {len(text)} characters in total]"
    return text


def list_files(folder: str = "") -> str:
    """List the files and folders in the workspace or one of its subfolders.

    Args:
        folder: Optional subfolder name. Leave empty for the top level.
    """
    base = _safe_path(folder)
    if not base.is_dir():
        return f"Folder not found: {folder}"
    lines = []
    for p in sorted(base.rglob("*"))[:100]:
        if p.name.startswith("."):
            continue
        rel = p.relative_to(WORKSPACE.resolve())
        lines.append(f"{rel}/" if p.is_dir() else f"{rel} ({p.stat().st_size} bytes)")
    return "\n".join(lines) if lines else "The workspace is empty."


def read_file(filename: str) -> str:
    """Read a file from the workspace. Supports text, code, markdown, json, csv, pdf, docx and xlsx.

    Args:
        filename: Name or relative path of the file, for example notes.txt or reports/q3.xlsx.
    """
    p = _safe_path(filename)
    if not p.is_file():
        return f"File not found: {filename}. Use list_files to see what exists."
    ext = p.suffix.lower()

    if ext == ".pdf":
        from pypdf import PdfReader
        pages = [(page.extract_text() or "") for page in PdfReader(str(p)).pages]
        text = "\n\n".join(f"[page {i + 1}]\n{t}" for i, t in enumerate(pages))
        return _truncate(text) if text.strip() else "This PDF has no extractable text (it may be a scan)."

    if ext == ".docx":
        from docx import Document
        doc = Document(str(p))
        parts = [para.text for para in doc.paragraphs if para.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return _truncate("\n".join(parts))

    if ext in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(str(p), data_only=True, read_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"[sheet: {ws.title}]")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 200:
                    out.append("[...more rows]")
                    break
                out.append(" | ".join("" if c is None else str(c) for c in row))
        return _truncate("\n".join(out))

    if ext in TEXT_LIKE or ext == "":
        return _truncate(p.read_text(encoding="utf-8", errors="replace"))

    return f"I can't read {ext} files yet."


def write_file(filename: str, content: str, overwrite: bool = False) -> str:
    """Create a file in the workspace. Supports text, code, markdown, json, csv, docx and xlsx.
    For xlsx and csv, give the content as CSV text with one row per line.
    For docx, each line becomes a paragraph. An existing file is only replaced if overwrite is true.

    Args:
        filename: Name or relative path of the file, for example shopping.txt or budget.xlsx.
        content: The full content to write.
        overwrite: Set to true only when the user explicitly wants to replace an existing file.
    """
    p = _safe_path(filename)
    ext = p.suffix.lower()
    if p.exists():
        if not overwrite:
            return f"{filename} already exists. Ask the user whether to overwrite it or use another name."
        shutil.copy2(p, p.with_name(p.name + ".bak"))   # veiligheidskopie
    p.parent.mkdir(parents=True, exist_ok=True)

    if ext == ".docx":
        from docx import Document
        doc = Document()
        for line in content.splitlines():
            doc.add_paragraph(line)
        doc.save(str(p))
    elif ext == ".xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        for row in csv.reader(io.StringIO(content)):
            ws.append(row)
        wb.save(str(p))
    elif ext in TEXT_LIKE or ext == "":
        p.write_text(content, encoding="utf-8")
    else:
        return f"I can't write {ext} files yet."
    return f"Saved {filename}."


def append_file(filename: str, content: str) -> str:
    """Add text to the end of an existing text, markdown, csv or log file in the workspace.
    The file is created if it does not exist yet.

    Args:
        filename: Name or relative path of the file.
        content: The text to add.
    """
    p = _safe_path(filename)
    if p.suffix.lower() not in TEXT_LIKE and p.suffix != "":
        return "Appending only works on plain text files. Use write_file for docx or xlsx."
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("a", encoding="utf-8") as f:
        f.write(("\n" if p.exists() and p.stat().st_size else "") + content)
    return f"Added to {filename}."


# --- Gmail & Drive tools (Google APIs) ---
GOOGLE_SCOPES = [
    "https://www.googleapis.com/auth/gmail.readonly",
    "https://www.googleapis.com/auth/gmail.send",
    "https://www.googleapis.com/auth/drive.file",
    "https://www.googleapis.com/auth/drive.readonly",
]


def _google_creds():
    """OAuth2 credentials voor Gmail en Drive. Eerste keer moet je inloggen via browser."""
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from google.auth.transport.requests import Request

    creds = None
    token_file = Path.home() / ".jarvis_token.json"
    if token_file.exists():
        creds = Credentials.from_authorized_user_file(str(token_file), GOOGLE_SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            credentials_file = os.getenv("GOOGLE_CREDENTIALS", "credentials.json")
            flow = InstalledAppFlow.from_client_secrets_file(credentials_file, GOOGLE_SCOPES)
            creds = flow.run_local_server(port=0)
        token_file.write_text(creds.to_json())
    return creds


def gmail_list(query: str = "is:unread", max_results: int = 5) -> str:
    """Toon recente Gmail berichten die voldoen aan een zoekopdracht.

    Args:
        query: Gmail zoekquery, bijvoorbeeld 'is:unread' of 'from:nieuwsbrief'.
        max_results: Maximale hoeveelheid berichten (standaard 5).
    """
    try:
        from googleapiclient.discovery import build

        service = build("gmail", "v1", credentials=_google_creds())
        result = service.users().messages().list(userId="me", q=query, maxResults=max_results).execute()
        messages = result.get("messages", [])
        if not messages:
            return "Geen e-mails gevonden."
        lines = []
        for m in messages:
            msg = service.users().messages().get(userId="me", id=m["id"], format="metadata").execute()
            headers = {h["name"]: h["value"] for h in msg["payload"].get("headers", [])}
            lines.append(f"- {headers.get('Date','?')}: {headers.get('Subject','(geen onderwerp)')} (van {headers.get('From','?')})")
        return "\n".join(lines)
    except Exception as e:
        return f"Fout bij Gmail: {e}"


def gmail_send(to: str, subject: str, body: str) -> str:
    """Verstuur een e-mail via Gmail.

    Args:
        to: E-mailadres van de ontvanger.
        subject: Onderwerp van de e-mail.
        body: Inhoud van de e-mail.
    """
    try:
        import base64
        from email.mime.text import MIMEText
        from googleapiclient.discovery import build

        service = build("gmail", "v1", credentials=_google_creds())
        message = MIMEText(body)
        message["to"] = to
        message["subject"] = subject
        raw = base64.urlsafe_b64encode(message.as_bytes()).decode()
        service.users().messages().send(userId="me", body={"raw": raw}).execute()
        return f"E-mail verzonden naar {to}."
    except Exception as e:
        return f"Fout bij verzenden: {e}"


def drive_list(folder_id: str = "root", max_results: int = 10) -> str:
    """Toon bestanden in een Google Drive map.

    Args:
        folder_id: ID van de Drive map (standaard 'root').
        max_results: Maximale hoeveelheid bestanden (standaard 10).
    """
    try:
        from googleapiclient.discovery import build

        service = build("drive", "v3", credentials=_google_creds())
        result = service.files().list(
            q=f"'{folder_id}' in parents and trashed=false",
            pageSize=max_results,
            fields="files(id, name, mimeType, modifiedTime)"
        ).execute()
        files = result.get("files", [])
        if not files:
            return "Drive-map is leeg."
        lines = [f"- {f['name']} ({f['mimeType']})" for f in files]
        return "\n".join(lines)
    except Exception as e:
        return f"Fout bij Drive: {e}"


# --- Home Assistant & Umbrel tools (later toe te voegen) ---
# Voorbeelden van functies die je later kunt toevoegen:
#
# def ha_states(entity_id: str = "") -> str:
#     """Haal status op van Home Assistant entiteiten."""
#     ...
#
# def ha_service(domain: str, service: str, entity_id: str) -> str:
#     """Roep een Home Assistant service aan."""
#     ...
#
# def umbrel_ssh(command: str) -> str:
#     """Voer een commando uit op je Umbrel via SSH."""
#     ...
#
# Registreer ze dan in de TOOLS dictionary hieronder.

TOOLS = {f.__name__: f for f in (
    list_files, read_file, write_file, append_file,
    gmail_list, gmail_send, drive_list,
)}


# --- Het model: antwoord streamen, zin voor zin ---
def pop_sentences(buf: str, first: bool):
    """Splitst gestreamde tekst in complete zinnen. Geeft (klare_zinnen, rest) terug.
    Bij de allereerste zin mag hij ook bij een komma knippen, zodat je eerder iets hoort."""
    parts = re.split(r"(?<=[.!?])\s+|\n+", buf)
    ready = [p for p in parts[:-1] if p.strip()]
    rest = parts[-1]
    if first and not ready and len(rest) > 45:
        commas = [m for m in re.finditer(r",\s", rest) if m.start() > 20]
        if commas:
            m = commas[-1]
            ready = [rest[:m.start() + 1]]
            rest = rest[m.end():]
    return ready, rest


def trim_history():
    """Haalt de oudste berichten weg als het gesprek te lang wordt, zodat het model niet vol loopt."""
    def size(m):
        return len(m.get("content") or "")

    while len(history) > 3 and sum(size(m) for m in history) > MAX_HISTORY_CHARS:
        del history[1]
        while len(history) > 2 and history[1]["role"] != "user":
            del history[1]


def ask_stream(prompt: str, stop: threading.Event):
    """Geeft het antwoord zin voor zin terug zodra die klaar is, en voert bestandsacties uit."""
    trim_history()
    history.append({"role": "user", "content": prompt})

    for _ in range(MAX_TOOL_ROUNDS):
        content, buf, tool_calls, said = "", "", [], 0
        try:
            for chunk in ollama_client.chat(model=LLM_MODEL, messages=history, tools=list(TOOLS.values()),
                                            stream=True, keep_alive=KEEP_ALIVE,
                                            options={"num_ctx": NUM_CTX}):
                if stop.is_set():
                    return
                msg = chunk["message"]
                piece = msg.get("content") or ""
                if piece:
                    content += piece
                    buf += piece
                    ready, buf = pop_sentences(buf, said == 0)
                    for s in ready:
                        said += 1
                        yield s
                if msg.get("tool_calls"):
                    tool_calls.extend(msg["tool_calls"])
            if buf.strip():
                yield buf.strip()
        finally:
            # Loopt ook bij onderbreken: wat er al was gezegd blijft in het gesprek staan
            entry = {"role": "assistant", "content": content}
            if tool_calls and not stop.is_set():
                entry["tool_calls"] = tool_calls
            if content or "tool_calls" in entry:
                history.append(entry)

        if stop.is_set() or not tool_calls:
            return

        for call in tool_calls:
            name = call["function"]["name"]
            args = call["function"]["arguments"]
            try:
                result = TOOLS[name](**args) if name in TOOLS else f"Unknown tool: {name}"
            except Exception as e:
                result = f"Error: {e}"
            history.append({"role": "tool", "tool_name": name, "content": str(result)})

    yield "I'm afraid I got stuck on that one, Sir."


# --- Spreken ---
def clean_for_speech(text: str) -> str:
    text = re.sub(r"```.*?```", " ", text, flags=re.S)   # codeblokken niet voorlezen
    text = re.sub(r"[*_#`>]+", "", text)                 # markdown weghalen
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def speak(source) -> np.ndarray | None:
    """Spreekt uit. `source` is een tekst, of een functie die (stop) -> zinnen teruggeeft.
    Onderbreek je hem, dan stopt hij meteen en krijg je de audio van jouw onderbreking terug
    (die wordt daarna gewoon verwerkt). Anders geeft dit None."""
    stop = threading.Event()
    done = threading.Event()
    if isinstance(source, str):
        sentences = iter(re.split(r"(?<=[.!?])\s+", clean_for_speech(source)))
    else:
        sentences = source(stop)

    drain(mic_q)
    gen = speaker.begin()

    def producer():
        in_code = False
        first_line = True
        try:
            for s in sentences:
                if stop.is_set():
                    break
                n = s.count("```")
                if in_code:
                    if n % 2:
                        in_code = False
                    continue
                if n % 2:
                    in_code = True
                    s = s.split("```")[0]
                s = clean_for_speech(s)
                if not s:
                    continue
                samples, sr = kokoro.create(s, voice=voice, speed=SPEED, lang="en-gb")
                if stop.is_set():
                    break
                # Tekst en stem komen samen: print pas als de zin ook echt wordt uitgesproken
                print(("J.A.R.V.I.S.: " if first_line else "              ") + s)
                first_line = False
                speaker.add(resample(samples, sr, AUDIO_RATE), gen)
                speaker.add(np.zeros(int(AUDIO_RATE * PAUSE), dtype=np.float32), gen)
        except Exception as e:
            print(f"[fout] {e}")
        finally:
            done.set()
            close = getattr(sentences, "close", None)
            if close:
                close()

    worker = threading.Thread(target=producer, daemon=True)
    worker.start()

    interrupted = None
    loud = []

    while not (done.is_set() and speaker.pending() == 0):
        try:
            block = mic_q.get(timeout=0.03)
        except queue.Empty:
            continue
        rms = float(np.sqrt(np.mean(block ** 2)))
        loud = loud + [block] if rms > BARGE_THRESHOLD else []

        if BARGE_IN and len(loud) >= BARGE_BLOCKS:
            speaker.stop()
            stop.set()
            # Neem direct door, zodat het begin van je zin niet verloren gaat
            interrupted = _record(mic_q, loud)
            break

    stop.set()
    worker.join(timeout=5)   # wacht tot de zin-generator klaar is, zodat het gesprek in de juiste volgorde blijft
    if interrupted is None:
        time.sleep(0.15)   # laat de laatste audio uit de hardwarebuffer lopen

    return interrupted


if __name__ == "__main__":
    print("J.A.R.V.I.S. standby. Zeg 'Jarvis' om te starten.")
    print("Zeg 'go to sleep' om terug naar standby te gaan, 'goodbye' om te stoppen.")
    if BARGE_IN:
        print("Je kunt Jarvis onderbreken door gewoon te gaan praten.")

    # Start de web-interface (Jarvis_web.py) in een achtergrondproces
    web_proc = None
    try:
        web_proc = subprocess.Popen(
            [sys.executable, "Jarvis_web.py"],
            cwd=str(Path(__file__).parent),
            creationflags=subprocess.CREATE_NEW_CONSOLE if sys.platform == "win32" else 0,
        )
        print(f"Web-interface gestart op http://localhost:5000 (PID {web_proc.pid})")
    except Exception as e:
        print(f"Kon web-interface niet starten: {e}")

    active = False
    last_activity = 0.0
    pending = None   # audio van een onderbreking, wacht om verwerkt te worden

    try:
        while True:
            was_interrupt = pending is not None
            if was_interrupt:
                audio, pending = pending, None
            else:
                audio = listen()

            text = transcribe(audio)
            if not text:
                continue
            print(f"Jij: {text}")
            lowered = text.lower()

            # Korte onderbreking als "stop" of "wait": niet doorsturen naar het model
            if was_interrupt and len(lowered.split()) <= 3 and any(w in lowered for w in STOP_WORDS):
                continue

            # Automatisch terug naar standby na lange stilte
            if active and IDLE_TIMEOUT and time.time() - last_activity > IDLE_TIMEOUT:
                active = False

            # Standby: wacht op het wake word
            if not active:
                if not any(w in lowered for w in WAKE_WORDS):
                    continue
                active = True
                last_activity = time.time()

                # Haal het wake word weg; bleef er niets van betekenis over, dan alleen bevestigen
                remainder = lowered
                for w in WAKE_WORDS:
                    remainder = remainder.replace(w, "")
                remainder = re.sub(r"[^\w\s]", "", remainder).strip()
                if len(remainder.split()) < 2:
                    pending = speak("Yes, Sir?")
                    continue

            # Vanaf hier: gewone conversatie, geen wake word meer nodig
            if "goodbye" in lowered:
                speak("Goodbye, Sir.")
                break

            if any(s in lowered for s in SLEEP_WORDS):
                speak("Going on standby, Sir.")
                active = False
                continue

            pending = speak(lambda stop: ask_stream(text, stop))
            last_activity = time.time()
    finally:
        mic_stream.stop()
        speaker.stream.stop()
        if web_proc and web_proc.poll() is None:
            web_proc.terminate()
            print("Web-interface gestopt.")